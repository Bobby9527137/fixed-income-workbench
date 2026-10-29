# -*- coding: utf-8 -*-
"""
端到端自测
==========
模拟数据 → 转换 → 断言关键清洗效果；并回读输出文件校验（roundtrip）。
运行前请先：python scripts/make_mock_data.py 且 python scripts/transform.py --input data/原始数据
运行：python scripts/test_transform.py
"""
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import pandas as pd

from bondtable import (DEFAULT_OUT_DIR, PROJECT_ROOT, convert_dataframe,
                       load_config, load_tasks, make_credit_text, make_local_text)

RAW_DIR = os.path.join(PROJECT_ROOT, "data", "原始数据")


def check(cond, msg):
    print(("  ✅ " if cond else "  ❌ ") + msg)
    return bool(cond)


def _time_ok(v):
    """时间段内每个时间点均已补零（如 09:30-10:10）"""
    tokens = re.findall(r"\d{1,2}:\d{2}", str(v))
    return bool(tokens) and all(re.fullmatch(r"\d{2}:\d{2}", t) for t in tokens)


def main() -> int:
    config = load_config()
    tasks = {}
    for name in ("模拟数据_地方债.xlsx", "模拟数据_信用债.xlsx"):
        for bond, sheet, df in load_tasks(os.path.join(RAW_DIR, name), config):
            tasks[bond] = df
    print("【转换前自检】")
    ok = check(set(tasks) == {"地方债", "信用债"}, "债种识别：两个文件分别识别为 地方债 / 信用债")

    results = {}
    print("\n【通用清洗规则】")
    date_cols = {"地方债": ("发行日", "缴款日", "上市日"), "信用债": ("发行日", "缴款日", "公告日")}
    for bond in ("地方债", "信用债"):
        bt = config["bond_types"][bond]
        result, warnings = convert_dataframe(tasks[bond], config, bond)
        results[bond] = result
        ok &= check(list(result.columns) == bt["target_headers"],
                    f"{bond}：输出为 {len(bt['target_headers'])} 个目标表头且顺序正确")
        ok &= check(not warnings, f"{bond}：源列全部命中，无缺失列警告")
        for col in date_cols[bond]:
            ok &= check(result[col].map(lambda v: bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", v))).all(),
                        f"{bond}：{col} 全部统一为 YYYY-MM-DD")
        ok &= check(result["发行量(亿)"].map(lambda v: isinstance(v, (int, float))).all(),
                    f"{bond}：发行量为数值类型")

    print("\n【地方债 · 口径与脏格式用例】")
    loc = results["地方债"]
    ok &= check(len(loc) == 11, f"末尾整行空行被自动剔除（源 12 行 → 输出 11 行，实际 {len(loc)} 行）")
    ok &= check("2024-06-14" in set(loc["发行日"]), "Excel 序列号 45457 → 2024-06-14")
    ok &= check("2024-06-10" in set(loc["发行日"]), "datetime 单元格（带 00:00:00）→ 2024-06-10")
    ok &= check("2024-06-17" in set(loc["上市日"]) and "2024.6.17" not in set(loc["上市日"]),
                "'2024.6.17' → 2024-06-17")
    ok &= check("2024-06-12" in set(loc["缴款日"]) and "2024年6月12日" not in set(loc["缴款日"]),
                "'2024年6月12日' → 2024-06-12")
    ok &= check((loc["债券代码"] == "2371157.IB").any() and (loc["债券代码"] == "2406995.IB").any(),
                "债券代码 '7位.IB' 格式按文本完整保留")
    ok &= check((loc["期限"] == "10Y").any() and (loc["期限"] == "3Y").any() and (loc["期限"] == "30Y").any(),
                "期限统一 Y 计年：'10年' → '10Y'、'3' → '3Y'、'30Y' 保持")
    ok &= check((loc["发行时间"] == "09:30-10:10").any(), "时间段 '9:30-10:10' → '09:30-10:10'（整段补零）")
    ok &= check(loc["发行时间"].map(_time_ok).all(), "发行时间段内每个时间点均已补零")
    ok &= check({"专项", "一般"} <= set(loc["专项"]), "专项←类型：地方债取值为'专项/一般'并正确带入")
    ok &= check({"新增", "置换", "再融资"} <= set(loc["性质"]), "性质取值（新增/置换/再融资）正常带入")
    ok &= check((loc["招标系统"] == "中央国债登记结算有限责任公司").all(),
                "招标系统取值'中央国债登记结算有限责任公司'正常带入")
    hl_kw = config["bond_types"]["地方债"].get("row_highlight", {}).get("keywords", [])
    hl_mask = loc["债券简称"].astype(str).map(lambda v: any(k in v for k in hl_kw))
    hl_pos = [i for i, m in enumerate(hl_mask) if m]
    ok &= check(len(loc) == 11, f"新增区域数据（含两只贵州债）后输出 11 行（实际 {len(loc)} 行）")
    ok &= check(len(hl_pos) == 5, "区域标黄：贵州×2 / 陕西 / 江西 / 西藏 5 行全部识别")
    ok &= check(list(loc["发行日"].astype(str)) == sorted(loc["发行日"].astype(str))
                and loc["债券简称"].iloc[0] == "24江苏债05" and loc["债券简称"].iloc[-1] == "24西藏债04",
                "整表按发行日升序排列（首行 2024-06-03 江苏、末行 2024-06-28 西藏）")
    ok &= check(hl_pos == [6, 7, 8, 9, 10],
                f"标黄 5 行不再置顶：位于各自发行日位置（第 7-11 行），第 1-6 行无标黄")
    ok &= check(list(loc.loc[loc["发行日"] == "2024-06-26", "债券简称"]) == ["24贵州债09", "24贵州债05"],
                "同日组内保持源表顺序：6/26 组 '24贵州债09' 在 '24贵州债05' 之前（不再按末尾数字排序）")

    print("\n【信用债 · 口径与脏格式用例（52 列真实导出表）】")
    cre = results["信用债"]
    ok &= check(len(cre) == 5, f"末尾整行空行被自动剔除（输出 5 行，实际 {len(cre)} 行）")
    ok &= check(cre["截标时间"].map(
                    lambda v: bool(re.fullmatch(r"\d{2}-\d{2} \d{2}:\d{2}", v))).all(),
                "截标时间全部统一为'月-日 时:分'（不带年份）")
    ok &= check("09-24 17:00" in set(cre["截标时间"]) and "2026-09-24 17:00" not in set(cre["截标时间"]),
                "截标时间 '2026-09-24 17:00' → '09-24 17:00'（源表完整时间自动去年份）")
    ok &= check("09-25 17:30" in set(cre["截标时间"]),
                "截标时间 '9-25 17:30' → '09-25 17:30'（未补零自动补齐）")
    ok &= check("09-26 17:00" in set(cre["截标时间"]),
                "截标时间为 datetime 单元格（读出 '2026-09-26 17:00:00'）→ '09-26 17:00'")
    ok &= check("10-09 17:00" in set(cre["截标时间"]),
                "截标时间 '10-09 17:00' 已符合口径，原样保持")
    ok &= check("2026-09-25" in set(cre["发行日"]) and "2026-09-26" in set(cre["发行日"]),
                "发行日 ← 发行起始日：'2026/9/25'→2026-09-25、Excel 序列号→2026-09-26")
    ok &= check((cre["发行量(亿)"] == 12.0).any(),
                "发行量：实际发行优先（计划 10 / 实际 12 → 取 12）")
    ok &= check((cre["发行量(亿)"] == 15.5).any() and (cre["发行量(亿)"] == 20.0).any(),
                "发行量：实际发行为空则回落计划发行（20 / '15.5' / 12.5 / 30 均正确带入）")
    ok &= check("2.5-3.2" in set(cre["申购区间"]) and (cre["申购区间"].astype(str) == "").sum() >= 3,
                "申购区间 ← 投标区间(%)：有值直取（'2.5-3.2'）；无值留空不占位")
    ok &= check("AAA/" in set(cre["主/债评级"]) and "AA+/AA+" in set(cre["主/债评级"])
                and "AAA/AAA" in set(cre["主/债评级"]),
                "主/债评级 ← 主体评级+债项评级 合成：'AAA/'（债项空）、'AA+/AA+'、'AAA/AAA'")
    ok &= check("4" in set(cre["YY评级"].astype(str)) and "" in set(cre["YY评级"].astype(str)),
                "YY评级 ← YY评分：数值带入；空值留空")
    ok &= check((cre["期限"] == "3Y").any() and (cre["期限"] == "5Y").any()
                and (cre["期限"] == "0.74Y").any() and (cre["期限"] == "5+NY").any()
                and (cre["期限"] == "1Y").any(),
                "期限 ← 发行期限 统一 Y 计年：'3Y' 保持、'5'→'5Y'、'0.74'→'0.74Y'、'5+N年'→'5+NY'、'1'→'1Y'")
    ok &= check("2026-09-30" in set(cre["缴款日"]) and "2026-09-22" in set(cre["公告日"]),
                "日期归一：'2026年9月30日'→2026-09-30、'2026/9/22'→2026-09-22")
    ok &= check("宁波银行股份有限公司" in set(cre["发行人"]) and "26宁波银行债01" in set(cre["债券简称"]),
                "发行人、债券简称直接带入")

    print("\n【地方债 · 一键文本播报】")
    txt = make_local_text(results["地方债"])
    expected_txt = (
        "*本周地方债待发行*【820.5500亿】\n"
        "\n6/3\n江苏150.0000亿（10Y）\n"
        "\n6/10\n四川66.6000亿（5Y）\n"
        "\n6/14\n广东88.5000亿（7Y）\n"
        "\n6/17\n浙江40.0000亿（3Y）\n"
        "\n6/24\n山东120.3500亿（30Y）\n"
        "\n6/25\n湖南55.5000亿（20Y）\n"
        "\n6/26\n贵州135.8000亿（7Y 10Y）\n"
        "\n6/27\n陕西75.2000亿（7Y）\n"
        "\n6/28\n江西58.6000亿（5Y）\n西藏30.0000亿（3Y）"
    )
    ok &= check(txt == expected_txt,
                "播报文本与预期逐字一致（头行双星号、4 位小数、地区短名、按日 M/D 分组）")
    ok &= check(txt.splitlines()[0] == "*本周地方债待发行*【820.5500亿】"
                and txt.index("6/3") < txt.index("6/25") < txt.index("6/28"),
                "播报文本：合计头行正确且按发行日升序分组（与表格排序口径一致）")
    ok &= check("贵州135.8000亿（7Y 10Y）" in txt and "贵州45" not in txt and "贵州90" not in txt,
                "同日同地区合并：6/26 两只贵州券合并一行（金额求和 135.80、期限 7Y 10Y 升序去重）")
    ok &= check(make_local_text(results["地方债"], word="今日").splitlines()[0]
                == "*今日地方债待发行*【820.5500亿】",
                "播报头行可切换：word='今日' 生效")

    print("\n【信用债（商金债） · 一键文本播报】")
    cre_txt = make_credit_text(results["信用债"])
    expected_cre = (
        "*本周商金债待发行*【90亿】\n"
        "\n9/24\n26宁波银行债01 20亿（3Y）\n"
        "\n9/25\n26杭城投MTN002 15.5亿（5Y）\n"
        "\n9/26\n26苏交通SCP003 12亿（0.74Y）\n"
        "\n10/8\n26杭银永续债01 12.5亿（5+NY）\n"
        "\n10/9\n26中石化SCP004 30亿（1Y）"
    )
    ok &= check(cre_txt == expected_cre,
                "商金债播报与预期逐字一致（头行双星号、每行=债券简称+金额（期限）、金额去尾零）")
    ok &= check(cre_txt.splitlines()[0] == "*本周商金债待发行*【90亿】"
                and cre_txt.index("9/24") < cre_txt.index("10/8") < cre_txt.index("10/9"),
                "商金债播报：按发行日升序分组（M/D 不补零，9 月与 10 月混排正常）")
    ok &= check(make_credit_text(results["信用债"], word="今日").splitlines()[0]
                == "*今日商金债待发行*【90亿】",
                "商金债播报头行可切换：word='今日' 生效")

    print("\n【输出文件回读校验】")
    from bondtable import write_output
    for bond in ("地方债", "信用债"):
        # 受控写出→回读→删除：不依赖 data/输出 里的历史/其他来源文件（真实联调会混入用户数据）
        self_out = os.path.join(DEFAULT_OUT_DIR, f"_selftest_{bond}_回读.xlsx")
        if os.path.exists(self_out):
            os.remove(self_out)
        out_path = write_output(results[bond], self_out)
        ok &= check(bool(out_path), f"{bond}：受控输出写入 data/输出/ 成功")
        back = pd.read_excel(out_path)
        ok &= check(list(back.columns) == config["bond_types"][bond]["target_headers"],
                    f"{bond}：输出文件回读列头一致")
        if bond == "地方债":
            from openpyxl import load_workbook
            ws = load_workbook(out_path).active
            fills = [str(ws.cell(row=r, column=1).fill.start_color.rgb)
                     for r in range(2, ws.max_row + 1)]
            yellow_pos = [i for i, f_ in enumerate(fills) if f_.endswith("FFFF00")]
            ok &= check(yellow_pos == [6, 7, 8, 9, 10],
                        f"整行标黄：第 7-11 行在各自发行日位置上黄色填充、其余行无填充、"
                        f"全表共 {len(yellow_pos)} 行标黄（应为 5）")
        os.remove(out_path)

    print("\n" + ("🎉 全部测试通过" if ok else "⚠️ 存在未通过的测试项"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
