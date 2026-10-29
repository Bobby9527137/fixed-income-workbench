# -*- coding: utf-8 -*-
"""
境内债制表工具 — 核心转换引擎
================================
职责：读取原始 Excel → 识别债券品种 → 按该品种的映射配置取列 → 清洗格式 → 输出标准表格
     （分品种配置：地方债 35 列源表 → 11 列目标；信用债 11 列源表 → 11 列目标）

设计原则：
- 映射关系全部来自 JSON 配置，调整表头不用改代码
- 纯 Python 跨平台（Windows / macOS / Linux），仅依赖 pandas + openpyxl
- 出错信息面向非技术用户
"""

import json
import os
import re
from datetime import date, datetime, time as dtime, timedelta

import pandas as pd

# ---------------------------------------------------------------------------
# 路径与常量
# ---------------------------------------------------------------------------
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFIG_PATH = os.path.join(PROJECT_ROOT, "config", "表头映射.json")
DEFAULT_OUT_DIR = os.path.join(PROJECT_ROOT, "data", "输出")

BOND_TYPES = ("地方债", "信用债")

_EXCEL_EPOCH = datetime(1899, 12, 30)  # Excel 日期序列号纪元（含 1900 闰年 bug 修正）


# ---------------------------------------------------------------------------
# 配置
# ---------------------------------------------------------------------------
def load_config(config_path: str = CONFIG_PATH) -> dict:
    """读取表头映射配置文件"""
    if not os.path.exists(config_path):
        raise FileNotFoundError(f"找不到映射配置文件：{config_path}")
    with open(config_path, "r", encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# 表头匹配
# ---------------------------------------------------------------------------
def _norm_header(h) -> str:
    """去掉表头里的空白字符（空格/换行/全角空格），用于稳健匹配真实文件中的折行表头"""
    return re.sub(r"[\s\u3000]+", "", str(h))


def _column_lookup(df: pd.DataFrame) -> dict:
    """规范化表头 → 实际列名（重名列时取第一个）"""
    lookup = {}
    for c in df.columns:
        k = _norm_header(c)
        if k not in lookup:
            lookup[k] = c
    return lookup


# ---------------------------------------------------------------------------
# 清洗函数
# ---------------------------------------------------------------------------
def _is_missing(v) -> bool:
    """空值判断：None / NaN / 空白字符串"""
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == ""
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def normalize_date(value) -> str:
    """各种日期写法 → 'YYYY-MM-DD' 文本。

    支持：datetime/date、'2024-06-05'、'2024/6/5'、'2024.6.5'、'2024年6月5日'、
    '2024-06-05 00:00:00'、Excel 日期序列号（如 45457）。
    无法识别时原样返回，便于人工核查。
    """
    if _is_missing(value):
        return ""
    if isinstance(value, (datetime, date)):
        d = value.date() if isinstance(value, datetime) else value
        return d.strftime("%Y-%m-%d")
    s = str(value).strip()
    s = re.sub(r"\s+00:00:00(\.0)?$", "", s)  # '2024-06-05 00:00:00' → 日期部分
    m = re.match(r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?$", s)
    if m:
        y, mo, dd = (int(x) for x in m.groups())
        return f"{y:04d}-{mo:02d}-{dd:02d}"
    m = re.match(r"^(\d{5})(\.0)?$", s)  # Excel 序列号（4~6 万覆盖 2009–2064 年）
    if m and 40000 <= int(m.group(1)) <= 60000:
        return (_EXCEL_EPOCH + timedelta(days=int(m.group(1)))).strftime("%Y-%m-%d")
    return s


def normalize_time(value) -> str:
    """时间补零：'9:30' → '09:30'；支持时间段 '9:30-10:10' → '09:30-10:10'。

    对单元格内出现的每个时间点分别补零（分钟亦补齐两位），'9:30:00' 的秒部分丢弃，
    时间段之间的连接符（-、—、~、至 等）原样保留。
    """
    if _is_missing(value):
        return ""
    if isinstance(value, (dtime, datetime)):
        return value.strftime("%H:%M")
    s = str(value).strip()
    dm = re.match(r"^\d{4}[-/.年]\d{1,2}[-/.月]\d{1,2}日?\s*", s)  # 单元格误存完整日期时间时，先剥离日期部分
    if dm:
        s = s[dm.end():]

    def _pad(m):
        return f"{int(m.group(1)):02d}:{int(m.group(2)):02d}"

    return re.sub(r"(\d{1,2}):(\d{1,2})(?::\d{1,2})?", _pad, s)


def normalize_deadline(value) -> str:
    """截标时间统一为'月-日 时:分'（2026-10-09 用户确认：不带年份，如 09-24 17:00）。

    兼容 '9-24 17:00'（补零）、'2026-09-24 17:00'（去年份）、datetime 单元格、
    '2026-9-24 17:00:00'（读表后字符串带秒）等写法；无时间部分时仅保留'月-日'。
    """
    if _is_missing(value):
        return ""
    if isinstance(value, (datetime, date)):
        base = value.strftime("%m-%d")
        if isinstance(value, datetime) and (value.hour or value.minute):
            return f"{base} {value.strftime('%H:%M')}"
        return base
    s = str(value).strip()
    m = re.search(r"(?:(\d{4})[-/.年])?(\d{1,2})[-/.月](\d{1,2})日?(?:\s*(\d{1,2}):(\d{1,2}))?", s)
    if not m:
        return s
    mo, dd = int(m.group(2)), int(m.group(3))
    if not (1 <= mo <= 12 and 1 <= dd <= 31):
        return s
    out = f"{mo:02d}-{dd:02d}"
    if m.group(4) is not None:
        out += f" {int(m.group(4)):02d}:{int(m.group(5)):02d}"
    return out


def compose_deadline(row) -> str:
    """信用债截标时间 = (招标日, 招标时间) 合成 → '月-日 时:分'（不带年份，2026-10-09 确认）。

    日期归一后取'月-日'，时间补零；招标时间为时间段时整段保留（'MM-DD HH:MM-HH:MM'）。
    """
    d_raw = row[0] if row else None
    t_raw = row[1] if len(row) > 1 else None
    if _is_missing(d_raw):
        return ""
    ds = str(normalize_date(d_raw))
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", ds)
    base = f"{m.group(2)}-{m.group(3)}" if m else ds
    ts = "" if _is_missing(t_raw) else str(normalize_time(t_raw)).strip()
    return f"{base} {ts}" if ts else base


def compose_range(row) -> str:
    """信用债申购区间 = (投标下限(%), 投标上限(%)) 合成 → '下限-上限'。

    任一为空则整格留空（2026-10-09 确认：有则有，没有则空，不用占位）。
    """
    low = row[0] if row else None
    high = row[1] if len(row) > 1 else None
    if _is_missing(low) or _is_missing(high):
        return ""
    return f"{str(low).strip()}-{str(high).strip()}"


def compose_rating(row) -> str:
    """信用债主/债评级 = (主体评级, 债项评级) 合成 → '主体/债项'（2026-10-09 确认示例 'AAA/'）。

    两项都有 → 'AA+/AA+'；债项为空 → 'AAA/'；两项都空 → 留空。
    """
    main = row[0] if row else None
    sub = row[1] if len(row) > 1 else None
    m_miss, s_miss = _is_missing(main), _is_missing(sub)
    if m_miss and s_miss:
        return ""
    return f"{'' if m_miss else str(main).strip()}/{'' if s_miss else str(sub).strip()}"


def compose_first(row):
    """多列取第一个非空值（如 发行量 = 实际发行(亿) 优先，空则取计划发行(亿)，2026-10-09 确认）。"""
    for v in row:
        if not _is_missing(v):
            return v
    return ""


def keep_text(value) -> str:
    """按文本保留：去空白、去掉数字的 '.0' 尾巴（保住源数据里的前导零）"""
    if _is_missing(value):
        return ""
    s = str(value).strip()
    if re.fullmatch(r"-?\d+\.0+", s):  # 12480034.0 → 12480034
        s = s.split(".", 1)[0]
    return s


def normalize_period(value) -> str:
    """期限统一为 Y 计年格式（2026-10-09 按用户口径确认）：
    '5Y' 保持、'10年' → '10Y'、纯数字 '3' → '3Y'、'5+N年' → '5+NY'；
    其他无法识别的写法原样保留，便于人工核查。
    """
    if _is_missing(value):
        return ""
    s = str(value).strip().replace("年", "Y")
    if s.lower().endswith("y"):
        return s[:-1] + "Y"
    if re.fullmatch(r"\d+(\.\d+)?", s):
        return f"{s}Y"
    return s


def format_amount(value, digits=None):
    """数值列：空 → ''；数值四舍五入保留 digits 位（None=保留原精度——
    精度交给写出环节的单元格格式与播报合计，2026-10-09 单日样本暴露的分层修复）；
    非数值原样保留便于人工核查"""
    if _is_missing(value):
        return ""
    try:
        f = float(str(value).replace(",", ""))
        return round(f, digits) if digits is not None else f
    except ValueError:
        return value


TRANSFORMS = {
    "date_normalize": normalize_date,
    "time_normalize": normalize_time,
    "deadline_normalize": normalize_deadline,
    "deadline_compose": compose_deadline,
    "range_compose": compose_range,
    "rating_compose": compose_rating,
    "first_nonempty": compose_first,
    "keep_text": keep_text,
    "period_y": normalize_period,
}


# ---------------------------------------------------------------------------
# 转换
# ---------------------------------------------------------------------------
def get_bond_config(config: dict, bond_type) -> dict:
    """取出某债券品种的映射配置段（bond_types.<品种>）"""
    sections = config.get("bond_types", {})
    if bond_type not in sections:
        raise ValueError(f"未配置品种【{bond_type}】的映射规则，请检查 config/表头映射.json")
    return sections[bond_type]


def convert_dataframe(df: pd.DataFrame, config: dict, bond_type):
    """原始 DataFrame → 按品种映射后的目标 DataFrame。

    返回 (结果表, 警告列表)
    """
    bt = get_bond_config(config, bond_type)
    mapping = bt["column_mapping"]
    lookup = _column_lookup(df)
    warnings = []
    out = pd.DataFrame(index=df.index)
    for col in bt["target_headers"]:
        item = mapping.get(col, {})
        if item.get("mode") == "placeholder_empty":  # 占位列（如信用债的主/债评级、发行人、YY评级、公告日）
            out[col] = ""
            continue
        src_names = item.get("sources") or ([item["source"]] if item.get("source") else [])
        src_norms = [_norm_header(s) for s in src_names]
        if any(s not in lookup for s in src_norms):
            missing = "、".join(n for n, s in zip(src_names, src_norms) if s not in lookup)
            warnings.append(f"原始表中未找到列【{missing}】，目标列【{col}】整列留空，请核对表头")
            out[col] = ""
            continue
        fn = TRANSFORMS.get(item.get("transform"))
        if len(src_norms) > 1:  # 多列合成（如 截标时间 = 招标日 + 招标时间）
            cols = [df[lookup[s]] for s in src_norms]
            out[col] = [fn(tuple(v)) if fn is not None else tuple(v) for v in zip(*cols)]
        else:
            series = df[lookup[src_norms[0]]]
            out[col] = series.map(fn) if fn is not None else series
    if "发行量(亿)" in out.columns:
        # 保留原始精度（播报合计需要，如 90.1234）；表格 2 位小数由写出环节单元格格式负责
        out["发行量(亿)"] = out["发行量(亿)"].map(format_amount)
    out = out.fillna("")
    for c in out.columns:  # 去除文本首尾空白（兼容 pandas 3.x 的字符串专用 dtype）
        if out[c].dtype == object or pd.api.types.is_string_dtype(out[c]):
            out[c] = out[c].map(lambda v: v.strip() if isinstance(v, str) else v)
    if len(out):  # 剔除整行全空的记录
        empty_row = out.apply(lambda r: all(str(v).strip() == "" for v in r), axis=1)
        out = out.loc[~empty_row]
    out = apply_row_order(out, bt)  # 整表排序（地方债按发行日升序，2026-10-09 v0.10）
    out = apply_row_highlight(out, bt)  # 行标黄（不置顶；attrs 携带标黄信息给写出环节）
    return out, warnings


def apply_row_order(df: pd.DataFrame, bt: dict) -> pd.DataFrame:
    """整表排序（2026-10-09 v0.10 新增，真实联调后确认）：bt["row_order"]["sort_by"]
    指定排序列时，整表按该列升序（stable：同值保持源表顺序；空值排最后）。
    当前用于地方债"按发行日排序"——标黄行不再置顶；未配置的品种（信用债）不排序。"""
    ro = bt.get("row_order") or {}
    col = ro.get("sort_by")
    if not col or col not in df.columns or df.empty:
        return df
    key = df[col].astype(str).str.strip()
    has = key.ne("") & key.ne("nan") & key.ne("NaT") & key.ne("None")
    order = sorted(range(len(df)), key=lambda i: (not bool(has.iloc[i]), key.iloc[i], i))
    return df.iloc[order].reset_index(drop=True)


def apply_row_highlight(df: pd.DataFrame, bt: dict) -> pd.DataFrame:
    """行标黄（2026-10-09 v0.10 修订：**只标黄、不再置顶**；行序由 row_order 决定）：
    配置了 row_highlight 的品种，匹配关键词的行整行标黄。
    标黄行位置与颜色写入 df.attrs["highlight_rows"] / ["highlight_color"]，
    由 write_output 落成整行黄色填充；未配置或无匹配时全部 False。"""
    hl = bt.get("row_highlight") or {}
    col, keywords = hl.get("match_column"), hl.get("keywords") or []
    ok = bool(col and keywords and col in df.columns and len(df))
    df.attrs["highlight_rows"] = ([any(k in v for k in keywords) for v in df[col].astype(str)]
                                  if ok else [False] * len(df))
    df.attrs["highlight_color"] = str(hl.get("fill_color") or "FFFF00")
    return df


# ---------------------------------------------------------------------------
# 读取与债种识别
# ---------------------------------------------------------------------------
BOND_ALIASES = {"地方债": ("地方政府债",), "信用债": ("商金债",)}


def detect_bond_type(name) -> "str | None":
    """按文件名 / 工作表名中的关键词识别债种；识别不到返回 None。
    2026-10-09 真实联调补充别名：'地方政府债 2026-09-22.xlsx' 文件名不含连续
    '地方债'三字 → 以'地方政府债'识别为地方债；'商金债'识别为信用债。"""
    s = str(name)
    for key in BOND_TYPES:
        if key in s:
            return key
    for key, aliases in BOND_ALIASES.items():
        if any(a in s for a in aliases):
            return key
    return None


def _find_header_row(raw: pd.DataFrame, source_headers: list, max_scan: int = 10) -> int:
    """在前 10 行里找与已知源表头命中最多的一行作为表头行（兼容表头上方有标题行的情况）"""
    known = {_norm_header(h) for h in source_headers}
    best_row, best_hits = 0, 0
    for i in range(min(max_scan, len(raw))):
        hits = sum(1 for v in raw.iloc[i].tolist() if _norm_header(v) in known)
        if hits > best_hits:
            best_row, best_hits = i, hits
    return best_row if best_hits >= 5 else 0


def read_sheet(xls: pd.ExcelFile, sheet_name: str, source_headers: list) -> pd.DataFrame:
    """读取工作表：自动定位表头行 → 全部数据按文本读取（防止前导零与格式丢失）"""
    raw = pd.read_excel(xls, sheet_name=sheet_name, header=None, dtype=str)
    if raw.shape[1] == 0:
        return pd.DataFrame()
    h = _find_header_row(raw, source_headers)
    header = [str(x).strip() for x in raw.iloc[h].tolist()]
    data = raw.iloc[h + 1:].reset_index(drop=True)
    data.columns = header
    return data


def load_tasks(source, config: dict, name: "str | None" = None):
    """解析一个 Excel 文件 → [(债种或 None, 工作表名, DataFrame)]。

    source：文件路径（命令行批处理）或二进制流（网页上传，2026-10-09 改为
    BytesIO 内存直读，不再落盘临时文件——旧版曾致项目根目录累积 tmp*.xlsx 泄漏）；
    name：流对应的文件名（用于债种识别与报错提示）。"""
    if isinstance(source, str):
        if not os.path.exists(source):
            raise FileNotFoundError(f"找不到文件：{source}")
        shown = os.path.basename(source)
    else:
        shown = str(name or getattr(source, "name", "") or "上传文件")
    try:
        xls = pd.ExcelFile(source)
    except Exception as e:
        raise ValueError(f"无法读取该文件（若为旧版 .xls 请先在 Excel 中另存为 .xlsx）：{e}") from e
    file_type = detect_bond_type(shown)
    # 未识别债种的工作表：用全部品种表头的并集定位表头行（品种可由用户在界面手动指定）
    union_headers: list = []
    for sec in config.get("bond_types", {}).values():
        union_headers.extend(sec.get("source_headers", []))
    tasks = []
    for sheet in xls.sheet_names:
        bond_type = detect_bond_type(sheet) or file_type
        headers = get_bond_config(config, bond_type)["source_headers"] if bond_type else union_headers
        try:
            df = read_sheet(xls, sheet, headers)
        except Exception as e:
            raise ValueError(f"读取工作表《{sheet}》出错：{e}") from e
        if df.empty:
            continue
        tasks.append((bond_type, sheet, df))
    return tasks


# ---------------------------------------------------------------------------
# 输出
# ---------------------------------------------------------------------------
def _display_width(s) -> int:
    """按显示宽度估算：中日韩字符按 2、其余按 1"""
    return sum(2 if ord(ch) > 127 else 1 for ch in str(s))


def write_output(df: pd.DataFrame, out_path: str, sheet_name: str = "制表结果") -> str:
    """写出带排版的 xlsx：表头加粗/填色/居中、全表细边框、自适应列宽、冻结首行。

    目标文件已存在时自动追加序号避免覆盖。返回实际写入的路径。
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    out_dir = os.path.dirname(out_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    base, ext = os.path.splitext(out_path)
    n = 1
    while os.path.exists(out_path):
        out_path = f"{base}_{n}{ext}"
        n += 1
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        head_font = Font(bold=True)
        head_fill = PatternFill("solid", fgColor="DDEBF7")
        hl_rows = {i for i, m in enumerate(df.attrs.get("highlight_rows") or []) if m}
        hl_fill = (PatternFill("solid", fgColor=str(df.attrs.get("highlight_color") or "FFFF00"))
                   if hl_rows else None)
        n_rows, n_cols = df.shape
        for j, col in enumerate(df.columns, start=1):
            hcell = ws.cell(row=1, column=j)
            hcell.font = head_font
            hcell.fill = head_fill
            hcell.border = border
            hcell.alignment = center
            width = _display_width(col) + 6
            is_amount_col = col == "发行量(亿)"
            for i in range(n_rows):
                c = ws.cell(row=i + 2, column=j)
                c.border = border
                if is_amount_col and isinstance(c.value, (int, float)):
                    c.number_format = "0.00"   # 表格口径：发行量显示 2 位小数（数据层保留原精度）
                if hl_fill is not None and i in hl_rows:  # 标黄行：整行黄色填充
                    c.fill = hl_fill
                if isinstance(c.value, str):
                    c.alignment = Alignment(vertical="center")
                else:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                width = max(width, _display_width("" if c.value is None else c.value) + 3)
            ws.column_dimensions[get_column_letter(j)].width = min(width, 42)
        ws.freeze_panes = "A2"
    return out_path


# ---------------------------------------------------------------------------
# 文本播报（2026-10-09 新增：网页"一键制作文本"用，地方债专用）
# ---------------------------------------------------------------------------
def region_short(short_name) -> str:
    """债券简称 → 地区短名（2026-10-09 真实数据联调确认的解析规则）：
    '26广东债47' → '广东'；'26内蒙古27'（无"债"字）→ '内蒙古'；
    '26青岛81' → '青岛'；'26广东债40(续发)' → '广东'（续发行并入本地区，不单列）。
    规则：去掉开头年份数字 → 取首个数字之前的部分 → 循环剥掉结尾的
    "续发"/"债"/"续"标记；无法解析时原样返回。"""
    s = re.sub(r"^\d+", "", str(short_name).strip())
    m = re.match(r"^[^\d]+", s)
    head = (m.group(0) if m else s).strip()
    while True:
        for suf in ("续发", "债", "续"):
            if head.endswith(suf) and len(head) > len(suf):
                head = head[: -len(suf)].strip()
                break
        else:
            break
    return head or str(short_name)


def _term_sort_key(term):
    """期限排序键：按数值升序（'5Y'→5、'0.74Y'→0.74、'5+NY'→5）；无法解析排最后。"""
    m = re.match(r"^(\d+(?:\.\d+)?)", str(term).strip())
    return (0, float(m.group(1))) if m else (1, 0.0)


def _fmt_amount(v) -> str:
    """金额去尾零：150.0 → '150'、'88.50' → '88.5'、120.35 → '120.35'。"""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    if pd.isna(f):
        return "0"
    return f"{f:g}"


def make_local_text(result, word: str = "本周") -> str:
    """地方债"待发行"播报文本（2026-10-09 按真实样例二次修正）：

        *本周地方债待发行*【532.3463亿】

        9/24
        江西261.5820亿（5Y 7Y 10Y 20Y 30Y）

    · 头行【】= 全表发行量合计，**固定 4 位小数**；word 可切换"本周/今日"；
    · 按发行日升序分组（组头 M/D 不补零），组间空一行；
    · 同日同地区**合并为一行**：地区短名 + 当日合计金额（4 位小数）
      +（当日各券期限，去重后按数值升序、空格分隔）；期限为空不带括号；
    · 文本与表格均按发行日升序（2026-10-09 v0.10 起表格标黄行不再置顶）。
    """
    if result is None or len(result) == 0:
        return ""
    df = pd.DataFrame(result).copy()
    df["__d"] = df["发行日"].astype(str).str.strip()
    df = df[df["__d"] != ""]
    if df.empty:
        return ""
    total = pd.to_numeric(df["发行量(亿)"], errors="coerce").fillna(0).sum()
    lines = [f"*{word}地方债待发行*【{total:.4f}亿】"]
    for d, grp in df.groupby("__d", sort=True):
        m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", d)
        head = f"{int(m.group(2))}/{int(m.group(3))}" if m else d
        lines.append("")
        lines.append(head)
        g2 = grp.copy()
        g2["__r"] = g2["债券简称"].map(region_short)
        for region, g in g2.groupby("__r", sort=False):  # sort=False 保持源表首次出现顺序
            amt = pd.to_numeric(g["发行量(亿)"], errors="coerce").fillna(0).sum()
            terms = sorted({str(t).strip() for t in g["期限"] if str(t).strip()},
                           key=_term_sort_key)
            per = f"（{' '.join(terms)}）" if terms else ""
            lines.append(f"{region}{amt:.4f}亿{per}")
    return "\n".join(lines)


def make_credit_text(result, word: str = "本周") -> str:
    """商金债（信用债）"待发行"播报文本（2026-10-09 按真实样例修正）：

        *本周商金债待发行*【20亿】

        9/24
        26重庆银行债01 20亿（3Y）

    · 头行【】= 全表发行量合计（去尾零）；word 可切换"本周/今日"；
    · 按发行日升序分组（组头 M/D 不补零），组间空一行；
    · 每行 = 债券简称 + 空格 + 发行量（去尾零）+（期限）；每券一行不合并；
      期限为空不带括号。
    """
    if result is None or len(result) == 0:
        return ""
    df = pd.DataFrame(result).copy()
    df["__d"] = df["发行日"].astype(str).str.strip()
    df = df[df["__d"] != ""]
    if df.empty:
        return ""
    total = pd.to_numeric(df["发行量(亿)"], errors="coerce").fillna(0).sum()
    lines = [f"*{word}商金债待发行*【{_fmt_amount(total)}亿】"]
    for d, grp in df.groupby("__d", sort=True):
        m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", d)
        head = f"{int(m.group(2))}/{int(m.group(3))}" if m else d
        lines.append("")
        lines.append(head)
        for _, r in grp.iterrows():
            per = f"（{r['期限']}）" if str(r["期限"]).strip() else ""
            lines.append(f"{str(r['债券简称']).strip()} {_fmt_amount(r['发行量(亿)'])}亿{per}")
    return "\n".join(lines)


def process_task(df: pd.DataFrame, bond_type, config: dict, out_dir: str = DEFAULT_OUT_DIR):
    """转换单个工作表并写出结果文件。返回 (输出路径, 结果表, 警告列表)"""
    result, warnings = convert_dataframe(df, config, bond_type)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    label = bond_type if bond_type else "未识别债种"
    out_path = os.path.join(out_dir, f"{label}_制表结果_{stamp}.xlsx")
    out_path = write_output(result, out_path)
    return out_path, result, warnings
