"""Excel / CSV 读写：不依赖 pandas，只用 openpyxl。"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any, Dict, List, Tuple

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except Exception:  # noqa: BLE001
    Workbook = load_workbook = None
    Font = PatternFill = Alignment = Border = Side = None
    get_column_letter = None


class ExcelError(ValueError):
    pass


def _require_openpyxl() -> None:
    if load_workbook is None or Workbook is None:
        raise ExcelError("缺少 openpyxl，请在工作台里重建该工具的环境")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _coerce(value: Any, kind: str) -> Any:
    text = _cell_text(value)
    if text == "":
        return None
    if kind == "number":
        try:
            return float(text.replace(",", "").replace("%", ""))
        except ValueError:
            return text
    if kind == "date":
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                continue
        return text
    return text


def read_table(
    filename: str,
    data: bytes,
    columns: List[dict],
) -> Tuple[List[Dict[str, Any]], List[str]]:
    """读 xlsx / csv 的第一张表，返回 (行列表, 警告)。"""
    suffix = (filename or "").lower().rsplit(".", 1)[-1]
    warnings: List[str] = []
    if suffix == "csv":
        rows = _read_csv(data, warnings)
    elif suffix in ("xlsx", "xlsm"):
        rows = _read_xlsx(data, warnings)
    else:
        raise ExcelError("只支持 .xlsx / .csv；旧版 .xls 请先另存为 .xlsx")
    if not rows:
        raise ExcelError("文件里没有读到数据行")

    header = [_cell_text(item) for item in rows[0]]
    mapping: Dict[int, dict] = {}
    for index, cell in enumerate(header):
        for column in columns:
            if cell in (column.get("label"), column.get("key")):
                mapping[index] = column
                break
    if not mapping:
        # 没匹配到表头时按列顺序对应
        for index, column in enumerate(columns[: len(header)]):
            mapping[index] = column
        warnings.append("没有匹配到表头，已按列顺序对应，请核对")

    result: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        if all(_cell_text(item) == "" for item in raw):
            continue
        item: Dict[str, Any] = {}
        for index, column in mapping.items():
            value = raw[index] if index < len(raw) else None
            item[column["key"]] = _coerce(value, column.get("type", "text"))
        result.append(item)
    return result, warnings


def _read_csv(data: bytes, warnings: List[str]) -> List[List[Any]]:
    text = None
    for encoding in ("utf-8-sig", "gbk", "utf-8"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ExcelError("CSV 编码无法识别，请另存为 UTF-8")
    reader = csv.reader(io.StringIO(text))
    return [row for row in reader]


def _read_xlsx(data: bytes, warnings: List[str]) -> List[List[Any]]:
    _require_openpyxl()
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return rows


def build_xlsx(sheets: List[dict]) -> bytes:
    """sheets = [{"name": "结果", "columns": [{"label": ...}], "rows": [[...]]}]"""
    _require_openpyxl()
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="EAF1F7")
    thin = Side(style="thin", color="D5DDE4")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_def in sheets:
        sheet = workbook.create_sheet(str(sheet_def.get("name") or "结果")[:31])
        columns = sheet_def.get("columns", [])
        label_row = [column.get("label", column.get("key", "")) for column in columns]
        sheet.append(label_row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        for row in sheet_def.get("rows", []):
            if isinstance(row, dict):
                values = [row.get(column.get("key")) for column in columns]
            else:
                values = list(row)
            sheet.append(values)
        for index, column in enumerate(columns, start=1):
            width = column.get("width")
            if not width:
                width = max(10, min(36, len(str(column.get("label", ""))) * 2 + 6))
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
