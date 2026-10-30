from __future__ import annotations

import base64
import importlib.util
import io
import unittest
from pathlib import Path

from openpyxl import load_workbook

from fi_toolbox import schema, selfcheck

ROOT = Path(__file__).resolve().parent.parent
APPS = (
    "bond_calculator",
    "repo_calculator",
    "cashflow_calendar",
    "portfolio_risk",
    "futures_basis",
    "bid_stats",
    "trade_extract",
)


def load_app(name: str):
    path = ROOT / "apps" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(f"test_app_{name}", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def number(text: str) -> float:
    return float(str(text).replace(",", ""))


class AppsTest(unittest.TestCase):
    def test_all_samples_and_selfcheck(self) -> None:
        for name in APPS:
            with self.subTest(app=name):
                module = load_app(name)
                data, errors = schema.normalize(module.FIELDS, module.sample(1))
                self.assertFalse(errors, errors)
                result = module.compute(data)
                self.assertTrue(result.get("metrics"))
                self.assertTrue(result.get("tables"))
                download = result.get("download") or {}
                self.assertTrue(download.get("base64"))
                raw = base64.b64decode(download["base64"])
                workbook = load_workbook(io.BytesIO(raw), read_only=True)
                self.assertTrue(workbook.sheetnames)
                workbook.close()
                check = selfcheck.run_checks(
                    tool_id=module.META["id"], meta=module.META, fields=module.FIELDS,
                    sample=module.sample, compute=module.compute,
                )
                self.assertTrue(check["ok"], check)

    def test_bond_calculator_values(self) -> None:
        module = load_app("bond_calculator")
        data, _ = schema.normalize(module.FIELDS, module.sample())
        result = module.compute(data)
        metrics = {item["label"]: item["value"] for item in result["metrics"]}
        self.assertAlmostEqual(number(metrics["净价"]), 101.7595, places=3)
        self.assertAlmostEqual(number(metrics["应计利息"]), 0.1233, places=3)
        self.assertAlmostEqual(number(metrics["修正久期"]), 8.7316, places=3)

    def test_repo_values(self) -> None:
        module = load_app("repo_calculator")
        data, _ = schema.normalize(module.FIELDS, module.sample())
        metrics = {item["label"]: item["value"] for item in module.compute(data)["metrics"]}
        self.assertEqual(metrics["实际天数"], "7")
        self.assertAlmostEqual(number(metrics["回购利息"]), 3.55, places=2)
        self.assertAlmostEqual(number(metrics["到期金额"]), 10003.55, places=2)

    def test_cashflow_values(self) -> None:
        module = load_app("cashflow_calendar")
        data, _ = schema.normalize(module.FIELDS, module.sample())
        metrics = {item["label"]: item["value"] for item in module.compute(data)["metrics"]}
        self.assertAlmostEqual(number(metrics["票息合计"]), 1245.50, places=2)
        self.assertAlmostEqual(number(metrics["本金兑付"]), 10000.0, places=2)

    def test_portfolio_values(self) -> None:
        module = load_app("portfolio_risk")
        data, _ = schema.normalize(module.FIELDS, module.sample())
        metrics = {item["label"]: item["value"] for item in module.compute(data)["metrics"]}
        self.assertAlmostEqual(number(metrics["总市值"]), 12500.0, places=2)
        self.assertAlmostEqual(number(metrics["总 DV01"]), 5.67, places=2)

    def test_futures_values(self) -> None:
        module = load_app("futures_basis")
        data, _ = schema.normalize(module.FIELDS, module.sample())
        metrics = {item["label"]: item["value"] for item in module.compute(data)["metrics"]}
        self.assertAlmostEqual(number(metrics["IRR"]), 2.5417, places=2)

    def test_bid_stats_values(self) -> None:
        module = load_app("bid_stats")
        data, _ = schema.normalize(module.FIELDS, module.sample())
        metrics = {item["label"]: item["value"] for item in module.compute(data)["metrics"]}
        self.assertAlmostEqual(number(metrics["投标总量"]), 9500.0, places=2)
        self.assertAlmostEqual(number(metrics["中标总量"]), 3500.0, places=2)

    def test_trade_extract_values(self) -> None:
        module = load_app("trade_extract")
        data, _ = schema.normalize(module.FIELDS, module.sample())
        result = module.compute(data)
        metrics = {item["label"]: item["value"] for item in result["metrics"]}
        self.assertEqual(metrics["识别成功"], "5")
        table = result["tables"][0]
        self.assertEqual(table["rows"][2]["name"], "26宁波银行债01")
        self.assertEqual(table["rows"][2]["counterparty"], "华信信托")


if __name__ == "__main__":
    unittest.main()
